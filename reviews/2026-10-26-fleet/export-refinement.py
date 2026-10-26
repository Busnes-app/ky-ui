#!/usr/bin/env python3
"""Export the separate refinement round as a verified workbook and offline gallery."""
import base64
import argparse
import csv
import hashlib
from html import escape
import json
import math
from pathlib import Path
import shutil
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output-dir', type=Path, default=ROOT / 'refinement-delivery')
OUT = parser.parse_args().output_dir.resolve()
OUT.mkdir(parents=True, exist_ok=True)
CSV_OUT = OUT / 'refinement-csvs' if OUT == ROOT else OUT
CSV_OUT.mkdir(exist_ok=True)


def read(name, default=None):
    path = ROOT / name
    return json.loads(path.read_text()) if path.exists() else default


def safe(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = value.rstrip()
        if value.startswith(('=', '+', '-', '@')):
            value = "'" + value
    return value


inputs = read('refinement-review-inputs.json')['files']
captures = read('refinement-captures.json')
score_record = read('refinement-scores.json')
ratings = score_record['files']
cohorts = read('refinement-cohorts.json')
sources = read('refinement-source-inputs.json')['repositories']
followup_captures = read('followup-captures-v2.json', read('followup-captures.json', []))
followup_review = read('followup-runtime-review.json')
runtime = read('refinement-runtime-review.json')
responses = read('refinement-review-responses.json', {})
if followup_review:
    responses.update({row['id']: {'disposition': row.get('disposition'), 'evidence': row.get('evidence')} for row in followup_review.get('findings', [])})
prs = read('refinement-pr-status.json', [])
progress = read('refinement-progress.json')
by_file = {row['file']: row for row in ratings}
assert set(by_file) == {row['file'] for row in inputs}
for row in captures + followup_captures:
    assert hashlib.sha256((ROOT / row['file']).read_bytes()).hexdigest() == row['sha256']
pairs = cohorts['matched_pairs']
assert len(pairs) * 2 == len(inputs)
for row in pairs:
    if row['product'] == 'KyNotes-server' and row['surface'] == 'appearance':
        row['comparison_limit'] = 'After has leftover Save/toast from capture order; before does not. Product regression unproven; delta confounded.'
    elif row['product'] == 'KyIdentity-server':
        row['comparison_limit'] = 'Both new fixtures load fonts. Prior baseline font404 correction is excluded from product improvement.'
    else:
        row['comparison_limit'] = 'Within-product comparison only; different products have different surfaces and capture modes.'

book = Workbook()
book.remove(book.active)
book.properties.title = 'Ky web refinement — matched screenshot evidence'
book.properties.creator = 'Claude Opus 5.5 council; Codex export verification'


def sheet(title, keys, rows):
    ws = book.create_sheet(title)
    ws.append(keys)
    for row in rows:
        ws.append([safe(row.get(key)) for key in keys])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for index, key in enumerate(keys, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = 65 if key in ('file', 'before', 'after', 'sha256', 'detail', 'observations', 'comparison_limit', 'evidence') else 24
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='182326')
    for cells in ws:
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(cell.value, str) and cell.value.startswith(('https://github.com/', 'refinement-products/', 'followup-products/')):
                cell.hyperlink = cell.value
                cell.style = 'Hyperlink'
    for index in range(2, ws.max_row + 1):
        ws.row_dimensions[index].height = 60
    with (CSV_OUT / (title.replace(' ', '-') + '.csv')).open('w', newline='', encoding='utf-8-sig') as handle:
        csv.writer(handle, lineterminator='\n').writerows(ws.values)
    return ws


sheet('Read first', ['topic', 'detail'], [
    {'topic': 'Result', 'detail': '9/10 target NOT achieved. Final image quality 6.1–8.0; Ky fit 6.8–8.0. Cohorts and products are not globally ranked.'},
    {'topic': 'Evidence', 'detail': '7 products, 216 staged design captures, 152 score rows / 146 exact-unique images, 76 matched pairs. 64 intermediate captures are not rated.'},
    {'topic': 'Follow-up evidence', 'detail': f'{len(followup_captures)} additional staged captures are separate from the scored round. Their numerical ratings are N/A. R1/R2 fixes and R3 disclosure require independent confirmation; see Follow-up review.'},
    {'topic': 'Method', 'detail': 'Claude Opus 5.5, ky-suite-v2. Weighted quality .20/.15/.20/.15/.15/.15; five-dimension Ky fit mean; Decimal ROUND_HALF_UP one decimal. Exact duplicates inherit scores.'},
    {'topic': 'Source/runtime review', 'detail': 'See separate Source review sheet. ' + ('Report available; read coverage and unresolved findings before acceptance.' if runtime else 'Independent review is PENDING; screenshots do not establish behavior gates.')},
    {'topic': 'Matched-state limitation', 'detail': 'Notes Appearance AFTER contains leftover Save/toast from capture order absent BEFORE; its negative delta is confounded and not a proven product regression.'},
    {'topic': 'Font fixture', 'detail': 'Identity before and after both use corrected fonts. Earlier font 404 baseline artifact is not a product improvement.'},
    {'topic': 'Capture modes', 'detail': 'Within-product matched fixtures. Chromium153/156,viewport versus full-page and tasks vary between products; compare within matched cohorts.'},
    {'topic': 'Scope and gaps', 'detail': 'This is a seven-product refinement round. Post/Calendar/Messages/Yard retain prior evidence. Drive Appearance and unlisted task surfaces were not newly captured.'},
    {'topic': 'Baseline', 'detail': 'Published baseline workbook and ratings remain unchanged. New scores must not silently replace that round.'},
])
sheet('Cohorts', ['product', 'cohort', 'task_view', 'phase', 'n', 'quality_mean', 'quality_min', 'quality_max', 'ky_fit_mean', 'ky_fit_min', 'ky_fit_max'], cohorts['cohort_product_phase'])
if progress:
    sheet('Progress', ['product', 'views', 'quality_before', 'quality_after', 'ky_fit_before', 'ky_fit_after', 'improvement'], progress['matched'])
    sheet('Earlier products', ['product', 'historical_start_range', 'latest_range'], progress['historical_context'])
sheet('Matched pairs', ['product', 'surface', 'viewport', 'theme', 'quality_before', 'quality_after', 'quality_delta', 'ky_fit_before', 'ky_fit_after', 'ky_fit_delta', 'byte_identical', 'comparison_limit', 'before', 'after'], pairs)
sheet('Ratings', ['file', 'product', 'phase', 'stage', 'observed_surface', 'cohort', 'theme', 'viewport', 'total', 'brand_total', 'scores', 'observations', 'ky_fit', 'confidence', 'artifacts', 'caveat', 'duplicate_of', 'sha256'], [dict(row, brand_total=row['ky_fit']['total']) for row in ratings])
sheet('Captures', ['file', 'product', 'surface', 'stage', 'phase', 'viewport', 'theme', 'width', 'height', 'source_revision', 'source_diff_sha256', 'driver', 'fixture', 'sha256', 'caveat'], captures)
sheet('PRs', ['repository', 'url', 'state', 'head', 'ci', 'review', 'review_caveat', 'observed_at'], prs)
if followup_captures:
    sheet('Follow-up captures', ['file', 'product', 'surface', 'stage', 'phase', 'viewport', 'theme', 'served_revision', 'checkout_revision', 'bundle_index_sha256', 'source_diff_sha256', 'sha256', 'rating', 'rating_reason', 'caveat'], followup_captures)
    sheet('Follow-up PRs', ['repository', 'url', 'state', 'head', 'ci', 'review', 'observed_at'], read('followup-pr-status.json', []))
    sheet('Follow-up review', ['section', 'detail'], [{'section': key, 'detail': value} for key, value in followup_review.items()] if followup_review else [{'section': 'Status', 'detail': 'R1 Identity and R2 Pulse implemented; R3 Notes disclosure/captures added with N/A limits. Independent confirmation is pending. No new numerical score.'}])
    sheet('Follow-up responses', ['id', 'disposition', 'response', 'verification'], [dict(value, id=key) for key, value in read('followup-review-responses.json', {}).items()])
if runtime:
    sheet('Source review', ['id', 'severity', 'origin', 'repository', 'path', 'summary', 'impact', 'fix', 'verification_result', 'current_disposition'], [dict(finding, current_disposition=responses.get(finding.get('id'), {})) for finding in runtime.get('findings', [])])
    sheet('Review coverage', ['repository', 'review_base', 'head', 'changed_paths', 'by_category', 'non_evidence_files'], runtime.get('coverage', []))
    sheet('Review limits', ['section', 'detail'], [{'section': key, 'detail': runtime.get(key)} for key in ('verdict', 'focus_results', 'executed_checks', 'prior_findings', 'reported_not_reexecuted', 'unknown')])
else:
    sheet('Source review', ['section', 'detail'], [{'section': 'Status', 'detail': 'Independent source/runtime confirmation pending; local implementation check claims are listed separately.'}])
checks = []
for filename in ('refinements-operations-response.json', 'refinement-notes-identity-response.json'):
    response = read(filename)
    for product, entry in response['repositories'].items():
        record = entry.get('checks', response.get('checks', {}).get(product, {}))
        for key, value in record.items():
            checks.append({'product': product, 'check': key, 'result': 'implementer record; see detail', 'detail': value})
sheet('Checks', ['product', 'check', 'result', 'detail'], checks)
gallery_sheet = book.create_sheet('Before and after')
for col in range(1, 15):
    gallery_sheet.column_dimensions[gallery_sheet.cell(1, col).column_letter].width = 9
cursor, embedded = 1, 0
for pair in pairs:
    gallery_sheet.cell(cursor, 1, f"{pair['product']} · {pair['surface']} · {pair['viewport']} · {pair['theme']} · Q {pair['quality_before']} → {pair['quality_after']}").font = Font(bold=True, size=12)
    max_height = 0
    for col, label, filename in [(1, 'Before', pair['before']), (8, 'After', pair['after'])]:
        gallery_sheet.cell(cursor + 1, col, label)
        picture = Image(ROOT / filename)
        ratio = min(1, 430 / picture.width)
        picture.width *= ratio
        picture.height *= ratio
        gallery_sheet.add_image(picture, gallery_sheet.cell(cursor + 2, col).coordinate)
        max_height = max(max_height, picture.height)
        embedded += 1
    cursor += math.ceil(max_height / 20) + 5
target = OUT / 'Ky-Web-Refinement-2026-10-26.xlsx'
book.save(target)
verified = load_workbook(target, read_only=True)
assert verified['Ratings'].max_row == 153 and verified['Captures'].max_row == 217
assert verified['Matched pairs'].max_row == 77
assert not any(cell.data_type == 'f' for ws in verified for row in ws for cell in row)
verified.close()
with zipfile.ZipFile(target) as archive:
    assert len([name for name in archive.namelist() if name.startswith('xl/media/')]) == embedded
if OUT != ROOT:
    shutil.copytree(ROOT / 'refinement-products', OUT / 'refinement-products', dirs_exist_ok=True)
if followup_captures and OUT != ROOT:
    shutil.copytree(ROOT / 'followup-products', OUT / 'followup-products', dirs_exist_ok=True)


def image(filename, label):
    data = base64.b64encode((ROOT / filename).read_bytes()).decode()
    return f'<figure><figcaption>{escape(label)}</figcaption><img loading="lazy" alt="{escape(label)}" src="data:image/png;base64,{data}"></figure>'


parts = ['''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Ky web refinement: matched evidence</title><style>
*{box-sizing:border-box}body{margin:0;color:#182326;background:#f8f6f0;font:1rem/1.5 system-ui,sans-serif}main{max-width:90rem;margin:auto;padding:1rem}h1{font-size:1.5rem}h2{font-size:1.25rem}p{max-width:85ch}summary,select{cursor:pointer;min-height:2.5rem}summary{font-weight:600;padding:1rem 0;border-top:1px solid #566461}select{font:inherit;color:#182326;background:white;border:1px solid #566461;border-radius:.5rem}label{display:inline-flex;align-items:center;gap:.5rem;margin:0 1rem 1rem 0}.pair{padding:1rem;background:white;border-radius:.5rem;margin-bottom:1rem}.pair h3{font-size:1.0625rem}.images{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:1rem}figure{margin:0;min-width:0}img{display:block;max-width:100%;height:auto;border:1px solid #566461}figcaption{font-weight:600;margin-bottom:.5rem}.meta{font-size:.875rem;color:#566461;overflow-wrap:anywhere}.pair[hidden]{display:none}summary:focus-visible,select:focus-visible{outline:2px solid #bf3f18;outline-offset:3px}@media(max-width:45rem){.images{grid-template-columns:1fr}}
</style><main><h1>Ky web refinement: matched evidence</h1><p>Claude Opus 5.5 · ky-suite-v2 · 76 matched views across seven products. Final screenshots score 6.1–8.0; the 9/10 target is not achieved. Screenshots do not establish behavior or accessibility acceptance.</p><p>Identity before/after both load the corrected fonts. Notes Appearance has a capture-state mismatch: its after image retains Save/toast absent before. Read those deltas as confounded. Other products, browsers and surfaces are not interchangeable ranking cohorts.</p><label>Size <select id="device"><option value="all">All</option><option value="desktop">Desktop</option><option value="mobile">Phone</option></select></label><label>Theme <select id="scheme"><option value="all">All</option><option value="light">Light</option><option value="dark">Dark</option></select></label>''']
for product in sorted({pair['product'] for pair in pairs}):
    group = [pair for pair in pairs if pair['product'] == product]
    parts.append(f'<details><summary>{escape(product)} · {len(group)} matched views</summary>')
    for pair in group:
        before, after = by_file[pair['before']], by_file[pair['after']]
        label = f"{pair['surface']} · {pair['viewport']} · {pair['theme']}"
        parts.append(f'<section class="pair" data-device="{escape(pair["viewport"])}" data-scheme="{escape(pair["theme"])}"><h3>{escape(label)}</h3><p>Quality {before["total"]} → {after["total"]}/10 · Ky fit {before["ky_fit"]["total"]} → {after["ky_fit"]["total"]}/10</p><div class="images">')
        parts.append(image(pair['before'], 'Before: ' + product + ' ' + label) + image(pair['after'], 'After: ' + product + ' ' + label))
        parts.append(f'</div><p class="meta">Before SHA256 {before["sha256"]}<br>After SHA256 {after["sha256"]}<br>{escape(pair["comparison_limit"])}</p></section>')
    parts.append('</details>')
if followup_captures:
    parts.append('<h2>Review follow-up stages · not numerically rated</h2><p>These supplement the scored round. Identity fixes an empty delegate menu; Pulse fixes Backup badge alignment; Notes documents broader secondary styling. Some Notes states are unreachable (N/A); the locked-session fixture is modelled by clearing storage except cookies. Independent confirmation status is in the workbook.</p>')
    for product in sorted({row['product'] for row in followup_captures}):
        parts.append(f'<details><summary>{escape(product)} · follow-up evidence</summary>')
        stages = sorted({row['stage'] for row in followup_captures if row['product'] == product})
        for stage in stages:
            parts.append(f'<details><summary>{escape(stage)}</summary>')
            for row in [r for r in followup_captures if r['product'] == product and r['stage'] == stage]:
                label = f"{row['surface']} · {row['viewport']} · {row['theme']}"
                device = 'mobile' if row['viewport'].startswith('390x') else 'desktop'
                scheme = 'dark' if 'dark' in row['file'] else 'light'
                parts.append(f'<section class="pair" data-device="{device}" data-scheme="{scheme}"><h3>{escape(label)}</h3>' + image(row['file'], label) + f'<p class="meta">Stage {escape(stage)} · score N/A (not rated)<br>Served revision {escape(str(row["served_revision"]))}<br>Bundle index SHA256 {escape(str(row["bundle_index_sha256"]))}<br>Image SHA256 {row["sha256"]}<br>{escape(row["caveat"])}</p></section>')
            parts.append('</details>')
        parts.append('</details>')
parts.append('''<p>Every screenshot is embedded; this page works offline. The workbook holds all 216 staged captures, 152 scores, 76 comparisons, PRs and check records. No new competitor comparison is implied.</p></main><script>const selects=[document.getElementById('device'),document.getElementById('scheme')];function filter(){for(const row of document.querySelectorAll('.pair'))row.hidden=selects.some(select=>select.value!=='all'&&row.dataset[select.id]!==select.value)}for(const select of selects)select.addEventListener('change',filter);</script></html>''')
page = OUT / 'refinement-gallery.html'
page.write_text(''.join(parts))
print(f'{target}:152scores,216captures,76pairs,{embedded} embedded images; reopened, no formulas')
print(f'{page}:{len(pairs)} matched pairs, {page.stat().st_size}bytes; all assets embedded')
