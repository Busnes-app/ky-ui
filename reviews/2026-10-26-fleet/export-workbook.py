#!/usr/bin/env python3
"""Export the fleet evidence snapshot. Requires openpyxl and Pillow, no network.

Run from any directory. Input records and screenshots live beside this script.
Missing scores remain blank; pending work never becomes a passing gate.
"""
import csv
import json
import math
import textwrap
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent


def read(name, default):
    path = ROOT / name
    return json.loads(path.read_text()) if path.exists() else default


def safe(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = value.rstrip()
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


fleet = read('fleet.json', {})
products = fleet.get('products', [])
captures = read('captures.json', [])
ratings = read('release-scores.json', {}).get('files', [])
checks = read('checks.json', [])
prs = read('prs.json', [])
gates = read('behavior-gates.json', [])
cohorts = read('cohort-results.json', [])
responses = read('review-responses.json', {})
findings = []
for review_file in ['runtime-source-review.json', 'runtime-source-review-v2.json', 'runtime-source-review-v3.json']:
    for finding in read(review_file, {}).get('findings', []) + read(review_file, {}).get('new_findings', []):
        findings.append({**finding, 'round': review_file, 'root_response': responses.get(finding.get('id'), {})})
book = Workbook()
book.remove(book.active)
book.properties.title = 'Ky web fleet — redesign evidence and PRs'
book.properties.creator = 'Codex and Claude Opus 5.5'
book.properties.description = 'Static evidence snapshot; 9/10 is a target, not certification.'


def sheet(title, columns, rows):
    ws = book.create_sheet(title)
    ws.append([label for _, label, _ in columns])
    for row in rows:
        ws.append([safe(row.get(key)) for key, _, _ in columns])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for index, (_, _, width) in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='182326')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        lines = 1
        for cell, (_, _, width) in zip(row, columns):
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            text = str(cell.value or '')
            lines = max(lines, sum(max(1, len(textwrap.wrap(line, max(8, int(width) - 2)))) for line in text.split('\n')))
            if text.startswith('https://github.com/'):
                cell.hyperlink = text
                cell.style = 'Hyperlink'
            elif text.startswith('products/') and (ROOT / text).is_file():
                cell.hyperlink = text
                cell.style = 'Hyperlink'
        ws.row_dimensions[row[0].row].height = min(300, max(30, lines * 15 + 6))
    with (ROOT / (title.replace(' ', '-') + '.csv')).open('w', newline='', encoding='utf-8-sig') as handle:
        csv.writer(handle, lineterminator='\n').writerows(ws.values)
    return ws


sheet('Read first', [('topic', 'Topic', 24), ('detail', 'Meaning', 110)], [
    {'topic': 'Status', 'detail': fleet.get('status', 'In progress')},
    {'topic': 'Scope', 'detail': fleet.get('scope', 'Ten active Ky web products plus server template')},
    {'topic': 'Scores', 'detail': 'Visual scores describe captured screens only. Missing scores are blank; cohort summaries are not whole-product rankings. No competitor superiority or 9/10 acceptance is implied.'},
    {'topic': 'Behavior', 'detail': 'Screenshot quality and behavior gates are separate. Untested gates prevent acceptance.'},
    {'topic': 'Capture conditions', 'detail': 'Synthetic fixtures; see individual source manifests and hashes. Viewport captures and full-page images are distinguished.'},
    {'topic': 'Merge and deployment', 'detail': 'Each PR needs passing CI and autonomous clearance at its current head. Nothing here authorizes merging or deployment.'},
])
sheet('Fleet', [
    ('repository', 'Repository', 24), ('status', 'Status', 30), ('base', 'Base revision', 44),
    ('theme_key', 'Retained storage key', 24), ('quality_score', 'Lowest captured cohort /10', 22),
    ('brand_fit', 'Lowest cohort Ky fit /10', 22), ('acceptance', 'Acceptance', 28), ('pr_url', 'PR', 64), ('notes', 'Scope / gaps', 70),
], products)
sheet('Captures', [(key, label, width) for key, label, width in [
    ('file', 'Capture', 70), ('product', 'Product', 24), ('stage', 'Stage', 10), ('phase', 'Before/after', 14),
    ('surface', 'Surface', 24), ('viewport', 'Viewport', 22), ('theme', 'Theme', 24),
    ('saved_theme', 'Saved theme', 20), ('os_scheme', 'OS scheme', 14), ('source_revision', 'Source revision', 44), ('sha256', 'SHA-256', 68), ('paired_file', 'Paired capture', 70), ('caveat', 'Capture caveat', 80),
]], captures)
sheet('Ratings', [
    ('file', 'Capture', 70), ('observed_product', 'Observed product', 24), ('observed_surface', 'Surface', 24),
    ('cohort', 'Cohort', 22), ('theme', 'Theme', 18), ('viewport', 'Viewport', 22), ('total', 'Quality /10', 14),
    ('ky_fit', 'Ky fit', 38), ('scores', 'Dimensions', 70), ('observations', 'Visible evidence', 110),
    ('confidence', 'Confidence', 16), ('state', 'Captured state', 16), ('artifacts', 'Capture artifacts', 50), ('duplicate_of', 'Duplicate: excluded from means', 70), ('unscored_reason', 'Unscored reason', 70), ('caveat', 'Limits', 100), ('sha256', 'SHA-256', 68),
], ratings)
sheet('Cohort results', [('product', 'Product', 24), ('cohort', 'Matched surface cohort', 24), ('n', 'Unique scored images', 18), ('mean', 'Quality mean /10', 18), ('brand_fit', 'Ky fit mean /10', 18), ('lowest_file', 'Lowest image /10', 18), ('comparison', 'Limits', 100)], cohorts)
sheet('Source review', [('round', 'Review round', 36), ('root_response', 'Response / current disposition', 100), ('id', 'Finding', 12), ('severity', 'Severity', 14), ('title', 'Finding', 78), ('verdict', 'Reviewer result', 20), ('regression', 'Introduced here', 18), ('fix', 'Suggested response', 110), ('verification', 'Verification', 100)], findings)
sheet('Behavior gates', [('product', 'Product', 24), ('gate', 'Gate', 24), ('result', 'Result', 18), ('revision', 'Revision', 44), ('evidence', 'Evidence', 100)], gates)
sheet('Checks', [('scope', 'Scope', 28), ('check', 'Check', 64), ('result', 'Result', 18), ('detail', 'Evidence / limits', 110), ('current_release_gate', 'Current release gate', 22)], checks)
sheet('PRs', [('repository', 'Repository', 26), ('url', 'PR', 68), ('state', 'State', 16), ('review_caveat', 'Review coverage caveat', 80), ('head', 'Checked head', 44), ('ci', 'CI', 20), ('review', 'Autonomous review', 26), ('merge_order', 'Merge order', 36), ('observed_at', 'Observed at (UTC)', 28)], prs)

gallery = book.create_sheet('Before and after')
for col in range(1, 15):
    gallery.column_dimensions[gallery.cell(1, col).column_letter].width = 9
row_index = 1
embedded = 0
for capture in captures:
    if capture.get('phase') != 'after' or not capture.get('paired_file') or capture.get('gallery_example') is not True:
        continue
    before, after = ROOT / capture['paired_file'], ROOT / capture['file']
    if not before.is_file() or not after.is_file():
        raise ValueError('Missing paired gallery input: ' + capture['file'])
    gallery.cell(row_index, 1, f"{capture['product']} · {capture['surface']} · {capture.get('theme', '')}").font = Font(bold=True, size=13)
    height = 0
    for col, label, path in [(1, 'Before', before), (8, 'After', after)]:
        gallery.cell(row_index + 1, col, label)
        image = Image(path)
        ratio = min(1, 430 / image.width)
        image.width *= ratio
        image.height *= ratio
        gallery.add_image(image, gallery.cell(row_index + 2, col).coordinate)
        height = max(height, image.height)
        embedded += 1
    row_index += math.ceil(height / 20) + 5

target = ROOT / 'Ky-Web-Fleet-2026-10-26.xlsx'
book.save(target)
with target.open('rb') as handle:
    delivered = load_workbook(handle, read_only=True)
    assert delivered['Fleet'].max_row == len(products) + 1
    assert delivered['Captures'].max_row == len(captures) + 1
    assert delivered['Ratings'].max_row == len(ratings) + 1
    assert delivered['PRs'].max_row == len(prs) + 1
    assert not any(cell.data_type == 'f' for ws in delivered for row in ws for cell in row)
    delivered.close()
with zipfile.ZipFile(target) as archive:
    assert len([name for name in archive.namelist() if name.startswith('xl/media/')]) == embedded
print(f'{target}: verified {len(products)} interfaces, {len(captures)} captures, {len(ratings)} scores, {len(prs)} PRs, {embedded} embedded images')
