#!/usr/bin/env python3
"""Export the council JSON and local screenshot evidence to XLSX + CSV.
Requires openpyxl 3.1.5 and Pillow. Run beside the council records; no network access.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone
import csv
import hashlib
import json
import math
import re
import textwrap

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DIMS = ['hierarchy', 'typography', 'density_layout', 'color', 'visible_affordances', 'coherence']
initial = json.loads((ROOT / 'initial-scores.json').read_text())
rerank = json.loads((ROOT / 'reranked-scores.json').read_text())
weights = rerank['rubric']['weights']
by_initial = {r['file']: r for r in initial['files']}
by_final = {r['file']: r for r in rerank['files']}
assert len(by_initial) == len(by_final) == 40 and by_initial.keys() == by_final.keys()
assert sum(Decimal(str(x)) for x in weights.values()) == 1
for row in by_final.values():
    score = by_final[row['duplicate_of']] if row.get('duplicate_of') else row
    total = sum(Decimal(str(weights[k])) * score['scores'][k] for k in DIMS)
    assert float(total.quantize(Decimal('.1'), rounding=ROUND_HALF_UP)) == score['total']

book = Workbook()
book.remove(book.active)
book.properties.title = 'Ky UI Council — ratings, redesigns and PRs'
book.properties.creator = 'Codex and Claude Opus 5.5'
book.properties.description = 'Static evidence snapshot. A 9/10 target is not an achieved score.'


def safe(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    # Score observations are data, including in Excel and CSV.
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def sheet(name, headers, rows, widths=None):
    ws = book.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([safe(x) for x in row])
    ws.freeze_panes = 'C2' if len(headers) > 6 else 'A2'
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='182326')
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Calibri', size=11, color='182326')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = '0.00'
    for n in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(n)].width = (widths or {}).get(n, 20)
    # Keep wrapped evidence readable in spreadsheet viewers that do not auto-fit on open.
    for row in ws.iter_rows(min_row=2):
        lines = max(sum(max(1, len(textwrap.wrap(part, width=max(8, int(ws.column_dimensions[cell.column_letter].width))))) for part in str(cell.value or '').split('\n')) for cell in row)
        ws.row_dimensions[row[0].row].height = min(390, max(20, lines * 15 + 4))
    if ws.max_row > 1:
        table = Table(displayName='T' + re.sub('[^a-zA-Z0-9]', '', name), ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.print_title_rows = '1:1'
    return ws


now = datetime.now(timezone.utc).isoformat(timespec='seconds')
release_summary = []
for p in sorted(ROOT.glob('implemented-*-release-scores.json')):
    review = json.loads(p.read_text())
    for product, result in review.get('product_visual_summary', {}).items():
        release_summary.append([product + ' — current visual result', f"Quality {result.get('equal_cohort_weight_mean')}/10; Ky brand fit {result.get('ky_fit_mean')}/10. Below the 9/10 target; see Acceptance and Behavior gates. Source: {p.name}"])
sheet('Start here', ['Item', 'Result / meaning'], [
    ['Generated UTC', now],
    ['Reviewers', 'Claude Opus 5.5: initial review, skill authoring, blind re-ranking, product implementation and visual review. Codex: orchestration, shared foundation, independent arithmetic checks and this export.'],
    ['Goal', 'Uniform Busnes cream/charcoal/orange identity; 9/10 overall is the target. No unverified 9/10 claim.'],
    ['Original inputs', '40 original PNGs, 39 unique. Private images remain outside the repository; this workbook embeds synthetic implementation screenshots only.'],
    ['How to compare', 'Use Cohort rankings for comparable task surfaces. Initial and improved-rubric scores describe unchanged originals. They are not redesign gains.'],
    ['Rounding', 'Per-image quality totals: one decimal. Improved-rubric cohort means: two decimals from unrounded weighted sums. Brand fit is separate.'],
    ['Observed baseline gaps', 'Most competitor comparisons mix theme, state or viewport; unmatched comparisons are flagged. Outlook captures are desktop clients of different generations.'],
    ['Missing evidence', 'Original set has no KyMessages. New KyMessages captures cover the console and authentication, not MemberHome or Element chat; absent cohorts are N/A, not zero.'],
    ['Owner correction', 'KyCalendar uses one left navigation region. The rejected double-left-panel layout was removed in stage 04.'],
    ['Stages', '01 Shell; 02 Task views; 03 Forms/dialogs; 04 Calendar single-sidebar correction; 05 usability refinements; 06 review corrections; 08 hover regressions. Test-only and DOM-identity fixes do not change screenshot pixels.'],
    ['Verification', 'Recorded tests are scoped evidence. Missing states, named-theme and full integration checks remain explicit; PR CI/review are separate gates.'],
    ['PRs', 'The PRs sheet records mergeable work and its live links when created. No merge or deployment is authorized by this report.'],
    ['Local image links', 'Stage evidence links are relative to this workbook. Download the screenshots folder beside it to use those links. The gallery embeds selected examples.'],
    ['Machine-readable sources', 'initial-scores.json and reranked-scores.json; implementation records, source-stage patches and verification.md are beside this workbook.'],
] + release_summary, {1: 26, 2: 125})

rank_rows = []
for cohort, value in initial['cohort_rankings'].items():
    for rank, r in enumerate(value['by_product'], 1):
        rank_rows.append(['Initial', cohort, rank, r['product'], r['n'], r['mean_total'], 'Uncontrolled originals', '', ', '.join(r['files'])])
for cohort, value in rerank['cohort_rankings'].items():
    for r in value['ranking']:
        rank_rows.append(['Improved rubric', cohort, r['rank'], r['product'], r['n'], r['mean'], 'Yes' if value['matched'] else 'No: ' + value.get('unmatched_reason', ''), (value['ky_gap'].get(r['product']) if isinstance(value.get('ky_gap'), dict) else value.get('ky_gap', 'N/A')), ', '.join(r['files'])])
sheet('Cohort rankings', ['Round', 'Cohort', 'Rank', 'Captured product', 'Unique images', 'Quality / 10', 'Matched state/theme/viewport', 'Ky gap to best competitor', 'Source filenames'], rank_rows, {2: 25, 4: 45, 7: 42, 9: 80})

score_rows, evidence_rows = [], []
for file in sorted(by_final):
    raw = by_final[file]
    r = {**by_final[raw['duplicate_of']], **raw} if raw.get('duplicate_of') else raw
    old = by_initial[file]
    fit = r.get('ky_fit', {})
    score_rows.append([file, r['observed_product'], r['observed_surface'], r['cohort'], r['state'], r['theme'], r['width'], r['height'], old['total'], r['total'], round(r['total'] - old['total'], 2), *[r['scores'][k] for k in DIMS], fit.get('value', 'N/A'), r.get('confidence', ''), raw.get('duplicate_of') or '', r.get('caveat', ''), r['sha256']])
    for dimension in DIMS:
        evidence_rows.append([file, r['observed_product'], dimension, weights[dimension], r['scores'][dimension], r['observations'][dimension], bool(raw.get('duplicate_of'))])
ws = sheet('Original scores', ['File', 'Product', 'Surface', 'Cohort', 'State', 'Theme', 'Width px', 'Height px', 'Initial quality', 'V2 quality', 'Method delta', *DIMS, 'Ky fit (separate)', 'Confidence', 'Duplicate of', 'Caveat', 'SHA-256'], score_rows, {1: 40, 2: 34, 3: 35, 4: 24, 6: 30, 20: 40, 21: 90, 22: 70})
ws.freeze_panes = 'E2'
for row in ws.iter_rows(min_row=2):
    row[8].number_format = row[9].number_format = '0.0'
sheet('Rating evidence', ['File', 'Product', 'Dimension', 'Weight', 'Score', 'Visible evidence', 'Inherited duplicate'], evidence_rows, {1: 40, 2: 34, 3: 25, 6: 110})
sheet('Rubric', ['Dimension / gate', 'Weight / requirement', 'Meaning'], [[k, weights[k], 'Fixed brand-neutral visual quality dimension'] for k in DIMS] + [
    ['Anchor 3', '', 'Broken: overlap, clipping, unreadable text or no primary action.'],
    ['Anchor 5', '', 'Works but fights you: competing emphasis, mixed styles, nested chrome or ambiguous states.'],
    ['Anchor 7', '', 'Solid: consistent components with noticeable minor slips.'],
    ['Anchor 9', '', 'Refined: clear emphasis, even rhythm and unambiguous states.'],
    ['Ky fit', 'Separate score', 'Palette, type/geometry, shell, accent discipline and cross-product cohesion. Never added to competitor quality.'],
    ['Keyboard', 'Pass / fail / untested', 'Keyboard task paths and dialog focus return.'],
    ['Focus', 'Pass / fail / untested', 'Visible focus, including valid native browser focus indication.'],
    ['Contrast', 'Pass / fail / untested', 'Measured text >=4.5:1, large text/control boundaries >=3:1.'],
    ['States', 'Pass / fail / untested', 'Empty/loading/error evidence and recovery.'],
    ['Responsive', 'Pass / fail / untested', '390/1280px, usable controls and contained overflow.'],
], {1: 28, 2: 30, 3: 115})
sheet('Score drift', ['File', 'Product', 'Old cohort', 'New cohort', 'Initial', 'Blind V2', 'Final V2', 'Delta', 'Cause', 'Explanation'], [[r['file'], r['observed_product'], r['prior_cohort'], r['cohort'], r['prior_total'], r['blind_total'], r['total'], r['delta'], r['cause'], r.get('explanation') or 'Below the 0.5 explanation threshold; see per-dimension evidence.'] for r in rerank['drift']['files']], {1: 40, 2: 34, 3: 22, 4: 24, 10: 115})

stage_rows = []
for after in sorted((ROOT / 'screenshots').glob('*-after-*.png')):
    match = re.match(r'(kypost|kycalendar|kymessages)-(\d\d)-after-(.*)\.png$', after.name)
    if not match:
        continue
    product, stage, surface = match.groups()
    before = after.with_name(after.name.replace('-after-', '-before-', 1))
    stage_rows.append([product, stage, surface, str(before.relative_to(ROOT)) if before.exists() else 'Not captured', str(after.relative_to(ROOT)), hashlib.sha256(after.read_bytes()).hexdigest(), 'Paired' if before.exists() else 'After only; unmatched'])
ws = sheet('Stage evidence', ['Product', 'Stage', 'Surface / viewport / theme', 'Before PNG', 'After PNG', 'After SHA-256', 'Pairing'], stage_rows, {3: 48, 4: 85, 5: 85, 6: 70, 7: 30})
for row in ws.iter_rows(min_row=2):
    for cell in (row[3], row[4]):
        if str(cell.value).startswith('screenshots/'):
            cell.hyperlink = cell.value
            cell.style = 'Hyperlink'

checks_file = ROOT / 'check-results.json'
checks = json.loads(checks_file.read_text()) if checks_file.exists() else [
    ['ky-ui', 'ESLint + unit', 'Pass', '8 unit tests; lint zero warnings', 'verification.md'],
    ['ky-ui', 'Chromium contracts', 'Pass', '20 checks; both Busnes themes and legacy hover', 'verification.md'],
    ['ky-ui', 'Selected consumer bytes', 'Pass', 'KyPost / KyCalendar / KyMessages match 0.3.0', 'verification.md'],
    ['Fleet', 'Public freshness', 'Pending rollout', 'Public branches still use the previous source copy', 'verification.md'],
    ['KyPost', 'Frontend suite', 'Pass; follow-up lint in progress', '1185 tests; typecheck/build/audit passed; lint follow-up handles baseline warnings', 'kypost-implementation.md'],
    ['KyCalendar', 'Local CI + browser', 'Pass; responsive follow-up in progress', '154 frontend tests, race/vet/fork/smoke; 20 browser checks', 'kycalendar-implementation.md'],
    ['KyMessages', 'Local CI + browser', 'Pass', '153 unit, race/vet; smoke 85 checks on free port; 7 Chromium/Firefox/captcha browser checks', 'kymessages-implementation.md'],
    ['KyMessages', 'MemberHome / Element', 'Untested', 'No real identity/chat integration capture', 'kymessages-implementation.md'],
    ['All', '9/10 acceptance', 'Not certified', 'Fresh visual review and behavior evidence must establish the result', 'verification.md'],
]
sheet('Checks', ['Product', 'Check / gate', 'Status', 'Evidence / limitation', 'Record'], checks, {1: 22, 2: 34, 3: 38, 4: 115, 5: 38})
prs_file = ROOT / 'pull-requests.json'
prs = json.loads(prs_file.read_text()) if prs_file.exists() else [{'repository': r, 'url': '', 'status': 'Preparing', 'base': b} for r, b in [('ky-ui', 'main'), ('kypost-server', 'main'), ('KyCalendar', 'main'), ('KyMessage-Server', 'master')]]
ws = sheet('PRs', ['Repository', 'PR URL', 'Base', 'Branch', 'Status', 'Merge / rollout note'], [[r['repository'], r.get('url', ''), r.get('base', ''), r.get('branch', 'ui/council-cohesion'), r.get('status', 'See live PR'), r.get('note', 'Human merge; no deployment implied.')] for r in prs], {1: 28, 2: 88, 4: 28, 5: 34, 6: 100})
for row in ws.iter_rows(min_row=2):
    if row[1].value:
        row[1].hyperlink = row[1].value
        row[1].style = 'Hyperlink'

rounds_file = ROOT / 'review-rounds.json'
round_status = json.loads(rounds_file.read_text()) if rounds_file.exists() else {}
acceptance, implementation_rows, implementation_evidence, gate_rows = [], [], [], []
for p in sorted(ROOT.glob('implemented-*-scores.json')):
    review = json.loads(p.read_text())
    round_name = review.get('round', p.stem)
    review_weights = review.get('rubric', {}).get('weights', weights)
    assert set(review_weights) == set(DIMS)
    assert all(Decimal(str(review_weights[k])) == Decimal(str(weights[k])) for k in DIMS), p.name
    for product, result in review.get('product_visual_summary', {}).items():
        acceptance.append([p.name, round_status.get(p.name, 'Recorded round; see source revision'), product, 'Product visual summary', result.get('equal_cohort_weight_mean'), result.get('ky_fit_mean'), result.get('files'), result.get('meets_9_quality_every_cohort', False), result.get('label', '')])
    for product, cohorts in review.get('cohort_aggregates', {}).items():
        for cohort, result in cohorts.items():
            acceptance.append([p.name, round_status.get(p.name, 'Recorded round; see source revision'), product, cohort, result.get('mean'), result.get('ky_fit_mean'), result.get('n'), (result.get('mean') or 0) >= 9, result.get('matched', 'Within-product visual evidence only')])
    for product, gates in review.get('gates', {}).items():
        for gate, result in gates.items():
            if isinstance(result, dict):
                gate_rows.append([p.name, product, gate, result.get('status', 'untested'), result.get('evidence', '')])
            else:
                gate_rows.append([p.name, product, gate, result, 'Review metadata; see source JSON.'])
    for result in review.get('files', []):
        scores = result.get('scores')
        if not isinstance(scores, dict):
            continue
        total = sum(Decimal(str(review_weights[k])) * scores[k] for k in DIMS)
        assert float(total.quantize(Decimal('.1'), rounding=ROUND_HALF_UP)) == result['total'], (p.name, result['file'])
        implementation_rows.append([p.name, round_name, result['file'], result.get('observed_product'), result.get('observed_surface'), result.get('cohort'), result.get('state'), result.get('theme'), result.get('viewport'), result['total'], result.get('ky_fit'), *[scores[k] for k in DIMS], result.get('capture'), result.get('sha256'), result.get('caveat')])
        for dimension in DIMS:
            implementation_evidence.append([p.name, result['file'], dimension, review_weights[dimension], scores[dimension], result.get('observations', {}).get(dimension, '')])
sheet('Acceptance', ['Record / review round', 'Evidence status', 'Product', 'Cohort / scope', 'Quality / 10', 'Ky fit (separate)', 'Scored captures', 'Scope meets 9', 'Comparison / limitation'], acceptance or [['Pending', 'Pending', 'All', 'Fresh review', None, None, 0, False, 'The 9/10 target is not yet certified.']], {1: 55, 2: 45, 3: 24, 4: 32, 9: 100})
ws = sheet('Implementation scores', ['Record', 'Round', 'File', 'Product', 'Surface', 'Cohort', 'State', 'Theme', 'Viewport', 'Quality / 10', 'Ky fit (separate)', *DIMS, 'Capture revision', 'SHA-256', 'Caveat'], implementation_rows, {1: 48, 2: 50, 3: 65, 4: 35, 18: 75, 19: 70, 20: 95})
for row in ws.iter_rows(min_row=2):
    row[9].number_format = row[10].number_format = '0.0'
sheet('Implementation evidence', ['Record', 'File', 'Dimension', 'Weight', 'Score', 'Visible evidence'], implementation_evidence, {1: 48, 2: 65, 3: 26, 6: 115})
sheet('Behavior gates', ['Record', 'Product', 'Gate', 'Status', 'Evidence / limitation'], gate_rows, {1: 55, 2: 26, 3: 26, 4: 24, 5: 125})

gallery = book.create_sheet('Screenshot gallery')
gallery.sheet_view.showGridLines = False
gallery.append(['Synthetic redesign screenshots — examples by stage; all captures are indexed in Stage evidence'])
gallery['A1'].font = Font(size=16, bold=True, color='182326')
for col in range(1, 16):
    gallery.column_dimensions[get_column_letter(col)].width = 11
pairs = [
    ('kypost', '01', 'inbox-desktop-light'), ('kypost', '02', 'inbox-selection-desktop-light'), ('kypost', '03', 'compose-desktop-light'),
    ('kycalendar', '01', 'calendar-desktop-light'), ('kycalendar', '02', 'week-desktop-light'), ('kycalendar', '03', 'new-event-mobile-dark'), ('kycalendar', '04', 'calendar-desktop-light'),
    ('kymessages', '01', 'overview-desktop-light'), ('kymessages', '02', 'audit-desktop-light'), ('kymessages', '03', 'settings-desktop-light'),
    ('kypost', '05', 'reader-mobile-light'), ('kypost', '05', 'contact-editor-mobile-dark'),
    ('kymessages', '05', 'audit-mobile-light'), ('kymessages', '05', 'health-desktop-dark'),
    ('kycalendar', '06', 'month-mobile-light'), ('kycalendar', '06', 'new-event-mobile-dark'),
    ('kymessages', '06', 'health-offline-mobile-light'),
    ('kycalendar', '08', 'nav-hover-desktop-light'), ('kymessages', '08', 'nav-hover-idle-mobile-light'),
]
row = 4
for product, stage, surface in pairs:
    files = [ROOT / 'screenshots' / f'{product}-{stage}-{state}-{surface}.png' for state in ('before', 'after')]
    if not all(p.exists() for p in files):
        continue
    gallery.cell(row, 1, f'{product} · stage {stage} · {surface}').font = Font(size=13, bold=True, color='182326')
    max_height = 0
    for col, label, path in zip((1, 8), ('Before', 'After'), files):
        cell = gallery.cell(row + 1, col, label + ' — open full PNG')
        cell.hyperlink = str(path.relative_to(ROOT))
        cell.style = 'Hyperlink'
        image = SheetImage(path)
        image.height = image.height * min(1, 430 / image.width)
        image.width = min(430, image.width)
        gallery.add_image(image, f'{get_column_letter(col)}{row + 2}')
        max_height = max(max_height, image.height)
    row += math.ceil(max_height / 20) + 5

out = ROOT / 'Ky-UI-Council-2026-10-26.xlsx'
book.save(out)
# Reopen with an independent parser path and validate delivered data, not just the generator.
with out.open('rb') as handle:
    delivered = load_workbook(handle, read_only=True, data_only=True)
    assert delivered['Original scores'].max_row == 41
    assert delivered['Rating evidence'].max_row == 241
    assert delivered['PRs'].max_row == len(prs) + 1
    assert {row[0] for row in delivered['Original scores'].iter_rows(min_row=2, values_only=True)} == set(by_final)
    actual = list(delivered['Original scores'].iter_rows(min_row=2, values_only=True))
    assert all(row[9] == ({**by_final[by_final[row[0]]['duplicate_of']], **by_final[row[0]]} if by_final[row[0]].get('duplicate_of') else by_final[row[0]])['total'] for row in actual)
    delivered.close()
for title in ('Cohort rankings', 'Original scores', 'Acceptance', 'Implementation scores', 'Behavior gates', 'Checks', 'PRs'):
    ws = book[title]
    with (ROOT / (title.replace(' ', '-') + '.csv')).open('w', newline='', encoding='utf-8-sig') as handle:
        csv.writer(handle, lineterminator='\n').writerows(ws.values)
print(f'{out}: verified {len(book.sheetnames)} sheets, 40 original rows, 240 dimension observations, {len(stage_rows)} stage images, {len(prs)} PR slots; {out.stat().st_size} bytes')
